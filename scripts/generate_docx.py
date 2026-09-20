#!/usr/bin/env python3
"""Local editable Word export from the same checklist JSON as HTML and PDF.

Word exports are teacher deliverables, never part of the site build. Text and
tables stay editable; sampling diagrams reuse the ReportLab drawing as an image.
"""
from __future__ import annotations

import argparse
import json
import re
import shutil
import subprocess
import tempfile
import uuid
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from lxml import etree
from reportlab.pdfgen import canvas

from generate_pdf import SamplingDiagram, register_fonts, validate

ROOT = Path(__file__).resolve().parent.parent
NAVY, INK, MUTED = '17334A', '253B49', '62737C'
SANS, SERIF, CJK = 'Source Sans 3', 'Source Serif 4', 'Noto Sans SC'
SUPPORTED = {'definition', 'answer-points', 'bullets', 'steps', 'comparison',
             'example', 'equivalence', 'callout', 'sampling-diagram'}


def set_font(style, name=SANS, size=10, color=INK, bold=False):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    style.font.italic = False
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attribute in list(fonts.attrib):
        if 'theme' in attribute.lower():
            del fonts.attrib[attribute]
    for script in ('ascii', 'hAnsi', 'cs'):
        fonts.set(qn(f'w:{script}'), name)
    fonts.set(qn('w:eastAsia'), CJK)


def runs(paragraph, text, bold=None):
    # Explicit CJK runs also work in Word installations using Latin-only defaults.
    for part in re.findall(r'[\u2e80-\ua4cf\uf900-\ufaff\uff00-\uffef]+|[^\u2e80-\ua4cf\uf900-\ufaff\uff00-\uffef]+', text):
        run = paragraph.add_run(part)
        if bold is not None:
            run.bold = bold
        if any(ord(c) >= 0x2e80 for c in part):
            run.font.name = CJK
            run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), CJK)


def paragraph(doc, text, style='Normal', **kwargs):
    p = doc.add_paragraph(style=style)
    runs(p, text, kwargs.pop('bold', None))
    for key, value in kwargs.items():
        setattr(p.paragraph_format, key, value)
    return p


def setup(doc, data, draft):
    s = doc.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    s.top_margin, s.bottom_margin = Mm(22), Mm(19)
    s.left_margin = s.right_margin = Mm(15)
    s.header_distance, s.footer_distance = Mm(11), Mm(10)
    styles = doc.styles
    set_font(styles['Normal'])
    styles['Normal'].paragraph_format.line_spacing = Pt(12.6)
    styles['Normal'].paragraph_format.space_after = Pt(3)
    styles['Normal'].paragraph_format.widow_control = True
    for name, font, size, color, bold in [
        ('Title', SERIF, 23, NAVY, False), ('Heading 1', SERIF, 24, NAVY, False),
        ('Heading 2', SANS, 10.7, NAVY, True), ('Subtitle', SANS, 9, MUTED, False),
        ('List Bullet', SANS, 10, INK, False), ('List Number', SANS, 10, INK, False),
        ('Caption', SANS, 8.5, MUTED, False), ('Header', SANS, 8, NAVY, False),
        ('Footer', SANS, 7.5, MUTED, False),
    ]:
        set_font(styles[name], font, size, color, bold)
    for name in ('Heading 1', 'Title'):
        # Remove borders inherited from Word's built-in Title template.
        for border in styles[name].element.xpath('./w:pPr/w:pBdr'):
            border.getparent().remove(border)
        f = styles[name].paragraph_format
        f.space_before, f.space_after = Pt(0), Pt(9)
        f.line_spacing = Pt(28)
        f.keep_with_next = True
    f = styles['Heading 2'].paragraph_format
    f.space_before, f.space_after, f.line_spacing = Pt(6), Pt(3), Pt(13)
    f.keep_with_next = True
    for name in ('List Bullet', 'List Number'):
        f = styles[name].paragraph_format
        f.left_indent, f.first_line_indent = Mm(4.5), Mm(-4.5)
        f.space_after, f.line_spacing = Pt(3), Pt(12.6)
    styles['Caption'].paragraph_format.space_after = Pt(3)
    styles['Caption'].paragraph_format.line_spacing = Pt(11)
    for name in ('Header', 'Footer'):
        styles[name].paragraph_format.tab_stops.clear_all()
    h = s.header.paragraphs[0]
    h.text = 'COMPUTER SCIENCE / 0478\tEXAM REVISION'
    h.paragraph_format.tab_stops.add_tab_stop(Mm(180), WD_TAB_ALIGNMENT.RIGHT)
    f = s.footer.paragraphs[0]
    status = 'CHAPTER DRAFT - LOCAL REVIEW' if draft else 'TEACHER WORD COPY'
    f.text = f'{status} | {data["label"]} | 2026-2028 | v{data["version"]}\tPage '
    f.paragraph_format.tab_stops.add_tab_stop(Mm(180), WD_TAB_ALIGNMENT.RIGHT)
    for instruction in ('PAGE', 'NUMPAGES'):
        if instruction == 'NUMPAGES':
            f.add_run(' of ')
        field = OxmlElement('w:fldSimple')
        field.set(qn('w:instr'), instruction)
        f._p.append(field)
    doc.core_properties.title = f'{data["label"]} {data["title"]}'
    doc.core_properties.subject = 'IGCSE Computer Science 0478 exam revision 2026-2028'
    doc.core_properties.author = ''
    doc.core_properties.keywords = 'teacher copy, local review' if draft else 'teacher copy'
    doc.core_properties.comments = ''


def table(doc, columns, rows):
    n = len(columns)
    fractions = [0.30, 0.70] if n == 2 else [1/n] * n
    t = doc.add_table(rows=1, cols=n)
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for col, fraction in zip(t.columns, fractions):
        col.width = Mm(180 * fraction)
    props = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        element = OxmlElement(f'w:{edge}')
        for k, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            element.set(qn(f'w:{k}'), value)
        borders.append(element)
    props.append(borders)
    margins = OxmlElement('w:tblCellMar')
    for edge, value in [('top', 45), ('bottom', 45), ('left', 100), ('right', 100)]:
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:w'), str(value)); el.set(qn('w:type'), 'dxa')
        margins.append(el)
    props.append(margins)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for index, row_data in enumerate([columns, *rows]):
        row = t.rows[0] if index == 0 else t.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for cell, text, fraction in zip(row.cells, row_data, fractions):
            cell.width = Mm(180 * fraction)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(11.1)
            runs(p, str(text), index == 0)
            for run in p.runs:
                run.font.size = Pt(9)
            if index == 0:
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'F3F6F7')
                cell._tc.get_or_add_tcPr().append(shade)
    return t


def sampling_image(block, temp):
    figure = SamplingDiagram(block)
    pdf = temp / f'{block["id"]}.pdf'
    c = canvas.Canvas(str(pdf), pagesize=(figure.width, figure.height))
    figure.drawOn(c, 0, 0)
    c.save()
    poppler = shutil.which('pdftoppm')
    if not poppler:
        raise RuntimeError('pdftoppm is required for sampling diagrams.')
    prefix = temp / block['id']
    subprocess.run([poppler, '-r', '240', '-singlefile', '-png', str(pdf), str(prefix)], check=True)
    return prefix.with_suffix('.png'), figure.height


def block(doc, data, temp):
    kind = data['type']
    paragraph(doc, data.get('title', data.get('term')), 'Heading 2')
    if kind in ('answer-points', 'bullets'):
        for text in data['items']:
            paragraph(doc, text, 'List Bullet')
    elif kind == 'steps':
        for index, text in enumerate(data['items'], 1):
            paragraph(doc, f'{index:02d}   {text}')
    elif kind == 'comparison':
        table(doc, data['columns'], data['rows'])
    elif kind in ('definition', 'callout'):
        paragraph(doc, data['text'])
    elif kind == 'example':
        paragraph(doc, data['problem'])
        for index, text in enumerate(data['steps'], 1):
            paragraph(doc, f'{index:02d}   {text}')
        paragraph(doc, data['result'], bold=True)
    elif kind == 'equivalence':
        table(doc, [item['label'] for item in data['values']], [[item['value'] for item in data['values']]])
    elif kind == 'sampling-diagram':
        path, _ = sampling_image(data, temp)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1
        pic = p.add_run().add_picture(str(path), width=Mm(180))
        pic._inline.docPr.set('descr', data['title'] + ' ' + '; '.join(
            f'{p["label"]}: {p["sampleRate"]} Hz, {p["resolution"]} bits per sample' for p in data['panels']))
    if data.get('hint'):
        paragraph(doc, data['hint'], 'Caption')


def embed_fonts(path, directory):
    """Embed the licensed full fonts using ECMA-376 font obfuscation."""
    specs = [(SANS, 'SourceSans3-Regular.ttf', 'Regular'),
             (SANS, 'SourceSans3-Semibold.ttf', 'Bold'),
             (SERIF, 'SourceSerif4-Regular.ttf', 'Regular'),
             (SERIF, 'SourceSerif4-Semibold.ttf', 'Bold'),
             (CJK, 'NotoSansSC-Regular.ttf', 'Regular')]
    with ZipFile(path) as archive:
        contents = {info.filename: archive.read(info.filename) for info in archive.infolist()}
    fonts = etree.fromstring(contents['word/fontTable.xml'])
    rel_ns = 'http://schemas.openxmlformats.org/package/2006/relationships'
    rel_file = 'word/_rels/fontTable.xml.rels'
    rels = etree.fromstring(contents[rel_file]) if rel_file in contents else etree.Element(f'{{{rel_ns}}}Relationships', nsmap={None: rel_ns})
    for i, (name, filename, face) in enumerate(specs):
        font = next((f for f in fonts if f.get(qn('w:name')) == name), None)
        if font is None:
            font = etree.SubElement(fonts, qn('w:font')); font.set(qn('w:name'), name)
        key = uuid.uuid5(uuid.NAMESPACE_URL, filename)
        value = bytearray((directory / filename).read_bytes())
        for j in range(32):
            value[j] ^= key.bytes[::-1][j % 16]
        target = f'fonts/checklist-{i}.odttf'
        contents['word/' + target] = bytes(value)
        rid = f'rIdChecklistFont{i}'
        etree.SubElement(rels, f'{{{rel_ns}}}Relationship', Id=rid, Target=target,
                         Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/font')
        embed = etree.SubElement(font, qn(f'w:embed{face}'))
        embed.set(qn('r:id'), rid); embed.set(qn('w:fontKey'), '{' + str(key).upper() + '}')
        embed.set(qn('w:subsetted'), 'false')
    contents['word/fontTable.xml'] = etree.tostring(fonts, xml_declaration=True, encoding='UTF-8', standalone=True)
    contents[rel_file] = etree.tostring(rels, xml_declaration=True, encoding='UTF-8', standalone=True)
    types = etree.fromstring(contents['[Content_Types].xml'])
    etree.SubElement(types, '{http://schemas.openxmlformats.org/package/2006/content-types}Default',
                     Extension='odttf', ContentType='application/vnd.openxmlformats-officedocument.obfuscatedFont')
    contents['[Content_Types].xml'] = etree.tostring(types, xml_declaration=True, encoding='UTF-8', standalone=True)
    with ZipFile(path, 'w', ZIP_DEFLATED) as archive:
        for filename, data in contents.items():
            archive.writestr(filename, data)


def render(data, output, fonts, draft=False):
    validate(data)
    unknown = {b['type'] for s in data['sections'] for b in s['blocks']} - SUPPORTED
    if unknown:
        raise ValueError(f'Add a reviewed editable Word layout for these blocks first: {sorted(unknown)}')
    # Even an explicit CLI path cannot accidentally put a teacher copy in public/.
    output = output.resolve()
    if not output.is_relative_to(ROOT / 'output/docx') or output.suffix.lower() != '.docx':
        raise ValueError('Word exports must be .docx files inside ignored output/docx/.')
    register_fonts(fonts)
    doc = Document()
    setup(doc, data, draft or data.get('sample', False))
    with tempfile.TemporaryDirectory(prefix='checklist_word_') as temp:
        page_titles = data.get('pdfPageTitles')
        last_page = None
        for index, section in enumerate(data['sections'], 1):
            for b_index, item in enumerate(section['blocks']):
                number = item['printPage'] if page_titles else index
                if number != last_page:
                    if last_page is None:
                        paragraph(doc, f'{data["label"]} {data["title"]}', 'Title')
                        paragraph(doc, 'IGCSE 0478 / 2026-2028 / Exam revision', 'Subtitle')
                    title = page_titles[number-1] if page_titles else section['title']
                    p = paragraph(doc, f'{number:02d} {title}', 'Heading 1')
                    p.paragraph_format.page_break_before = last_page is not None
                    last_page = number
                if b_index == 0 and section.get('summary'):
                    paragraph(doc, section['summary'])
                block(doc, item, Path(temp))
        output.parent.mkdir(parents=True, exist_ok=True)
        doc.save(output)
    embed_fonts(output, fonts)
    print(f'Generated local teacher Word copy: {output}')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--fonts', type=Path, default=ROOT / 'assets/fonts')
    parser.add_argument('--draft', action='store_true')
    args = parser.parse_args()
    render(json.loads(args.input.read_text()), args.output, args.fonts, args.draft)
